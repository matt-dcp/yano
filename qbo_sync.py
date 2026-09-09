#!/usr/bin/env python3
"""Casa Yano monthly QBO close: install the prior month's Yanonali YTD P&L.

Charity (books@drivencap.com) publishes monthly financials to the DCP Bookkeeping
shared drive, usually around the 2nd. The Yanonali YTD Profit and Loss carries
one column per month and supersedes every earlier QBO file we hold, so installing
it is a replace, not an append.

Run:  python3 qbo_sync.py [--month YYYY-MM] [--no-push] [--dry-run]

Exit codes: 0 ok or already current, 1 not published yet / precondition failed,
2 build or reconciliation failed.
"""
import argparse
import calendar
import json
import shutil
import subprocess
import sys
from datetime import date
from pathlib import Path

import openpyxl

SITE = Path(__file__).resolve().parent
DATA = SITE / "data"
ARCHIVE = DATA / "_superseded"
REPORTS = (Path.home() / "Library/CloudStorage/GoogleDrive-matt@drivencap.com"
           / "Shared drives/DCP Bookkeeping/Source Documents/DCP Wealth Fund"
           / "Draft Financial Reports")
TARGET = DATA / "QBO-Monthly-PL.xlsx"
# Yanonali is one class inside the DCP Wealth Fund books; this is the property-only P&L.
PATTERN = "*Yanonali-YTD Profit and Loss*.xlsx"


def prior_month(today: date) -> date:
    first = today.replace(day=1)
    return (first - __import__("datetime").timedelta(days=1)).replace(day=1)


def month_label(d: date) -> str:
    return f"{d:%b} {d.year}"        # "Aug 2026" — matches the xlsx sheet header


def js_month_label(d: date) -> str:
    # build.py normalises the sheet header to "Aug '26" in data.js. The two
    # formats must be mapped or the reconciliation silently looks in the wrong key.
    return f"{d:%b} '{d:%y}"


def month_columns(xlsx: Path) -> list[str]:
    ws = openpyxl.load_workbook(xlsx, data_only=True).active
    for row in ws.iter_rows(min_row=1, max_row=12, values_only=True):
        cells = [str(c).strip() for c in row if c is not None]
        if any(c.endswith(("2025", "2026", "2027")) and len(c) == 8 for c in cells):
            return [c for c in cells if c != "Total"]
    return []


def line_totals(xlsx: Path, label: str) -> dict:
    """Pull Rent / Total for Expenses / NOI for one month column."""
    ws = openpyxl.load_workbook(xlsx, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = col = None
    for i, row in enumerate(rows):
        vals = [str(c).strip() if c is not None else "" for c in row]
        if label in vals:
            hdr_i, col = i, vals.index(label)
            break
    if col is None:
        return {}
    out = {}
    for row in rows[hdr_i + 1:]:
        name = str(row[0]).strip() if row and row[0] is not None else ""
        if name in ("Rent", "Total for Expenses", "Net Operating Income"):
            try:
                out[name] = round(float(row[col]), 2)
            except (TypeError, ValueError):
                pass
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--month", help="YYYY-MM to install (default: prior month)")
    ap.add_argument("--no-push", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    want = (date(int(args.month[:4]), int(args.month[5:7]), 1)
            if args.month else prior_month(date.today()))
    label = month_label(want)
    print(f"Casa Yano QBO sync — looking for {label}")

    # already current?
    if TARGET.exists() and label in month_columns(TARGET):
        print(f"Already current: {TARGET.name} already contains {label}. Nothing to do.")
        return 0

    folder = REPORTS / f"{want:%B} {want.year}"
    if not folder.is_dir():
        print(f"NOT PUBLISHED YET: no folder {folder.name!r} on the shared drive.")
        return 1
    hits = sorted(folder.glob(PATTERN))
    if not hits:
        print(f"NOT PUBLISHED YET: {folder.name} exists but has no {PATTERN}.")
        return 1
    src = hits[-1]
    print(f"  found {src.name}")

    cols = month_columns(src)
    if label not in cols:
        print(f"INCOMPLETE: {src.name} has columns {cols} — no {label}.", file=sys.stderr)
        return 1
    expected = line_totals(src, label)
    if not expected:
        print(f"UNREADABLE: could not read {label} totals from {src.name}.", file=sys.stderr)
        return 2
    print(f"  source {label}: rent {expected.get('Rent'):,.2f} | "
          f"expenses {expected.get('Total for Expenses'):,.2f} | "
          f"NOI {expected.get('Net Operating Income'):,.2f}")

    if args.dry_run:
        print(f"  would install -> {TARGET}")
        return 0

    # ---- install ----------------------------------------------------------
    ARCHIVE.mkdir(exist_ok=True)
    if TARGET.exists():
        shutil.move(str(TARGET), ARCHIVE / f"QBO-Monthly-PL.pre-{label.replace(' ','')}.xlsx")
    # single-month QBO-PL-*.xlsx files are all contained in the YTD file now
    for old in DATA.glob("QBO-PL-*.xlsx"):
        shutil.move(str(old), ARCHIVE / old.name)
        print(f"  archived superseded {old.name}")
    shutil.copy2(src, TARGET)
    print(f"  installed {TARGET.name}")

    # ---- build ------------------------------------------------------------
    r = subprocess.run([sys.executable, "build.py"], cwd=SITE, capture_output=True, text=True)
    print(r.stdout[-1200:])
    if r.returncode != 0:
        print(r.stderr[-2000:], file=sys.stderr)
        return 2

    # ---- reconcile the build against the bookkeeper's statement -----------
    js = (SITE / "public" / "data.js").read_text()
    blob = json.loads(js[js.find("{"): js.rfind("}") + 1])
    q = blob.get("qboPL", {})
    jlabel = js_month_label(want)
    if jlabel not in q.get("months", []):
        print(f"RECONCILE FAILED: build did not pick up {jlabel} "
              f"(months present: {q.get('months', [])}).", file=sys.stderr)
        return 2
    got = {}
    for row in q.get("rows", []):
        if row.get("label") in expected:
            m = row.get("monthly") or {}
            v = m.get(jlabel) if isinstance(m, dict) else None
            if v is not None:
                got[row["label"]] = round(float(v), 2)
    bad = [k for k, v in expected.items() if abs(got.get(k, 0) - v) > 0.01]
    if bad:
        for k in bad:
            print(f"RECONCILE FAILED {k}: source {expected[k]:,.2f} vs build {got.get(k, 0):,.2f}",
                  file=sys.stderr)
        return 2
    print(f"  reconciled {label} to the penny: " +
          " | ".join(f"{k} {v:,.2f}" for k, v in expected.items()))

    pf = blob["proForma"]
    print(f"\n  FY pro forma: gross ${pf['gross']:,.0f} | NOI cash ${pf['noiCash']:,.0f} "
          f"| normalized ${pf['noiNormalized']:,.0f} | closed months {pf['closedMonths']}")

    # ---- ship -------------------------------------------------------------
    subprocess.run(["git", "add", "public/data.js"], cwd=SITE, check=True)
    if subprocess.run(["git", "diff", "--cached", "--quiet"], cwd=SITE).returncode == 0:
        print("data.js unchanged; nothing to commit.")
        return 0
    msg = (f"Update dashboard - {label} QBO close: NOI cash ${pf['noiCash']:,.0f}\n\n"
           f"Installed {src.name} as QBO-Monthly-PL.xlsx; it carries every month "
           f"through {label} and supersedes the prior QBO files.\n\n"
           f"Reconciled to the bookkeeper's statement to the penny:\n"
           f"  rent ${expected['Rent']:,.2f} | expenses "
           f"${expected['Total for Expenses']:,.2f} | NOI "
           f"${expected['Net Operating Income']:,.2f}\n\n"
           f"FY pro forma gross ${pf['gross']:,.0f}, NOI cash ${pf['noiCash']:,.0f}, "
           f"normalized ${pf['noiNormalized']:,.0f}.\n\n"
           "Co-Authored-By: Claude Opus 5 <noreply@anthropic.com>")
    subprocess.run(["git", "commit", "-m", msg], cwd=SITE, check=True)
    if args.no_push:
        print("Committed. --no-push set, not pushing.")
        return 0
    subprocess.run(["git", "push", "origin", "main"], cwd=SITE, check=True)
    print("Pushed. Vercel will deploy.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
